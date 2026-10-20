import openpyxl
from typing import Optional
from app.core.database import DatabaseManager

class ExcelImporter:
    def __init__(self, db: DatabaseManager):
        self.db = db

    def import_processes_excel(self, file_path: str, cobit_version: str) -> int:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        
        col_map = {str(h).strip().lower(): idx for idx, h in enumerate(headers) if h}
        required = ["domain", "process_code", "name"]
        for req in required:
            if req not in col_map:
                raise ValueError(f"Kolom wajib '{req}' tidak ditemukan dalam Excel")

        count = 0
        with self.db.get_connection() as conn:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row or not any(row):
                    continue
                domain = str(row[col_map["domain"]]).strip()
                code = str(row[col_map["process_code"]]).strip()
                name = str(row[col_map["name"]]).strip()
                desc = str(row[col_map.get("description", 3)] or "").strip() if "description" in col_map else ""

                conn.execute(
                    """INSERT INTO processes (cobit_version, domain, process_code, name, description)
                       VALUES (?, ?, ?, ?, ?)""",
                    (cobit_version, domain, code, name, desc)
                )
                count += 1
        return count
